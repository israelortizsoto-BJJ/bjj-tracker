"""Integration tests for Google Sheets publishing (mocked API)."""

from __future__ import annotations

import json
from pathlib import Path
from unittest.mock import MagicMock

import pytest

ROOT = Path(__file__).resolve().parent.parent
DATASETS = ROOT / "TEST_DATASETS"

import sys

sys.path.insert(0, str(ROOT))

from google_sheets_publisher import (  # noqa: E402
    GooglePublishConfig,
    GooglePublishOptions,
    build_batch_requests,
    load_google_config,
    publish_timeline_xlsx,
    _extract_sheet_payload,
)
from timeline_builder import build_timeline  # noqa: E402


class MockExecute:
    def __init__(self, payload: dict):
        self.payload = payload

    def execute(self):
        return self.payload


class MockSpreadsheets:
    def __init__(self, store: dict):
        self.store = store

    def create(self, body=None):
        spreadsheet_id = "new-sheet-123"
        tab_id = 1001
        self.store["spreadsheetId"] = spreadsheet_id
        self.store["sheets"] = [
            {"properties": {"sheetId": tab_id, "title": body["sheets"][0]["properties"]["title"]}}
        ]
        self.store["create_body"] = body
        return MockExecute(
            {
                "spreadsheetId": spreadsheet_id,
                "sheets": [{"properties": {"sheetId": tab_id, "title": body["sheets"][0]["properties"]["title"]}}],
            }
        )

    def get(self, spreadsheetId=None, fields=None):
        return MockExecute(
            {
                "sheets": self.store.get(
                    "sheets",
                    [{"properties": {"sheetId": 42, "title": "Timeline"}}],
                )
            }
        )

    def batchUpdate(self, spreadsheetId=None, body=None):
        self.store.setdefault("batch_updates", []).append(body)
        replies = []
        for request in body.get("requests", []):
            if "addSheet" in request:
                replies.append(
                    {"addSheet": {"properties": {"sheetId": 2002, "title": request["addSheet"]["properties"]["title"]}}}
                )
        return MockExecute({"replies": replies})

    def values(self):
        return self

    def update(self, spreadsheetId=None, range=None, valueInputOption=None, body=None):
        self.store["values_update"] = {
            "spreadsheetId": spreadsheetId,
            "range": range,
            "values": body.get("values"),
        }
        return MockExecute({})


class MockSheetsService:
    def __init__(self):
        self.store: dict = {}

    def spreadsheets(self):
        return MockSpreadsheets(self.store)


@pytest.fixture
def sample_xlsx(tmp_path):
    output = tmp_path / "Timeline.xlsx"
    build_timeline(
        DATASETS / "A_small_3projects_20tasks.csv",
        output,
        ROOT / "colors.yaml",
        report_path=tmp_path / "ValidationReport.json",
    )
    return output


@pytest.fixture
def google_config(tmp_path):
    credentials = tmp_path / "service_account.json"
    credentials.write_text(
        json.dumps({"type": "service_account", "client_email": "test@example.com"}),
        encoding="utf-8",
    )
    config_path = tmp_path / "google_config.yaml"
    config_path.write_text(
        f"credentials_path: {credentials}\n"
        "default_sheet_id: null\n"
        "publish:\n"
        "  tab_name: Timeline\n",
        encoding="utf-8",
    )
    return load_google_config(config_path), credentials


class TestGoogleSheetsPublisher:
    def test_create_new_sheet(self, sample_xlsx, google_config, tmp_path):
        config, _ = google_config
        mock = MockSheetsService()

        result = publish_timeline_xlsx(
            sample_xlsx,
            config,
            source_name="Q1 Schedule",
            service_factory=lambda _path: mock,
        )

        assert result.success is True
        assert result.status == "published"
        assert result.sheet_id == "new-sheet-123"
        assert "new-sheet-123" in result.sheet_url
        assert mock.store["values_update"]["range"] == "Timeline!A1"
        assert len(mock.store["values_update"]["values"]) > 3

    def test_update_existing_sheet(self, sample_xlsx, google_config):
        config, _ = google_config
        mock = MockSheetsService()
        mock.store["sheets"] = [{"properties": {"sheetId": 99, "title": "Timeline"}}]

        result = publish_timeline_xlsx(
            sample_xlsx,
            config,
            sheet_id="existing-sheet-456",
            service_factory=lambda _path: mock,
        )

        assert result.success is True
        assert result.sheet_id == "existing-sheet-456"
        assert mock.store["batch_updates"]
        delete_requests = mock.store["batch_updates"][0]["requests"]
        assert any("deleteSheet" in req for req in delete_requests)
        assert any("addSheet" in req for req in delete_requests)

    def test_publish_failure_returns_failed_result(self, sample_xlsx, google_config):
        config, _ = google_config

        def broken_factory(_path):
            raise RuntimeError("API unavailable")

        result = publish_timeline_xlsx(
            sample_xlsx,
            config,
            service_factory=broken_factory,
        )

        assert result.success is False
        assert result.status == "failed"
        assert "API unavailable" in result.error

    def test_extract_payload_includes_merges_and_freeze(self, sample_xlsx):
        from openpyxl import load_workbook

        ws = load_workbook(sample_xlsx)["Timeline"]
        options = GooglePublishOptions()
        payload = _extract_sheet_payload(ws, options)

        assert payload.values
        assert payload.merge_ranges
        assert payload.frozen_row_count == 3
        assert payload.frozen_column_count == 5
        assert payload.row_groups

    def test_batch_requests_include_formatting_and_groups(self, sample_xlsx):
        from openpyxl import load_workbook

        ws = load_workbook(sample_xlsx)["Timeline"]
        payload = _extract_sheet_payload(ws, GooglePublishOptions())
        requests = build_batch_requests(payload, sheet_id=123, options=GooglePublishOptions())

        assert any("mergeCells" in req for req in requests)
        assert any("repeatCell" in req for req in requests)
        assert any("updateSheetProperties" in req for req in requests)
        assert any("addDimensionGroup" in req for req in requests)


class TestGooglePublishCLI:
    def test_publish_failure_preserves_excel_output(self, tmp_path, google_config, capsys):
        config, _ = google_config
        missing_config = tmp_path / "missing_google.yaml"
        output = tmp_path / "Timeline.xlsx"
        report = tmp_path / "ValidationReport.json"

        from timeline_builder import main

        code = main(
            [
                str(DATASETS / "A_small_3projects_20tasks.csv"),
                "-o",
                str(output),
                "--report",
                str(report),
                "--publish-google",
                "--google-config",
                str(missing_config),
            ]
        )

        assert code == 0
        assert output.exists()
        assert report.exists()
        captured = capsys.readouterr()
        assert "Google Publish Status: failed" in captured.out

    def test_load_google_config_requires_credentials_path(self, tmp_path):
        config_path = tmp_path / "bad.yaml"
        config_path.write_text("default_sheet_id: null\n", encoding="utf-8")
        with pytest.raises(ValueError, match="credentials_path"):
            load_google_config(config_path)
