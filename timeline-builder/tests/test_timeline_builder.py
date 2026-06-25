"""Automated validation suite for Timeline Builder production hardening."""

from __future__ import annotations

from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
DATASETS = ROOT / "TEST_DATASETS"
CONFIG = ROOT / "colors.yaml"

import sys

sys.path.insert(0, str(ROOT))

from color_matcher import task_style  # noqa: E402
from models import INFO_COLUMNS, TaskRow  # noqa: E402
from timeline_builder import (  # noqa: E402
    build_timeline,
    build_week_columns,
    build_workbook,
    critical_week_index,
    group_tasks_by_project,
    load_config,
    read_schedule,
    weeks_for_range,
    week_start,
)
from validation import validate_and_normalize  # noqa: E402


@pytest.fixture
def config_bundle():
    return load_config(CONFIG)


@pytest.fixture
def small_input():
    return DATASETS / "A_small_3projects_20tasks.csv"


@pytest.fixture
def edge_input():
    return DATASETS / "D_edge_cases.csv"


@pytest.fixture
def built_small(small_input, tmp_path):
    output = tmp_path / "Timeline.xlsx"
    build_timeline(
        small_input,
        output,
        CONFIG,
        report_path=tmp_path / "ValidationReport.json",
    )
    return load_workbook(output)


class TestTimelineGeneration:
    def test_builds_from_small_dataset(self, small_input, tmp_path):
        output = tmp_path / "Timeline.xlsx"
        result = build_timeline(
            small_input,
            output,
            CONFIG,
            report_path=tmp_path / "ValidationReport.json",
        )
        assert result.output_path.exists()
        assert result.output_path.stat().st_size > 5000
        assert result.report.timeline_generated

    def test_builds_from_medium_dataset(self, tmp_path):
        output = tmp_path / "Timeline.xlsx"
        build_timeline(
            DATASETS / "B_medium_10projects_100tasks.csv",
            output,
            CONFIG,
            report_path=tmp_path / "ValidationReport.json",
        )
        wb = load_workbook(output)
        assert wb.active.max_row >= 100

    def test_builds_from_large_dataset(self, tmp_path):
        output = tmp_path / "Timeline.xlsx"
        build_timeline(
            DATASETS / "C_large_30projects_515tasks.csv",
            output,
            CONFIG,
            report_path=tmp_path / "ValidationReport.json",
        )
        assert output.exists()

    def test_builds_with_extra_columns(self, tmp_path):
        output = tmp_path / "Timeline.xlsx"
        build_timeline(
            DATASETS / "D_edge_cases_extra_columns.csv",
            output,
            CONFIG,
            report_path=tmp_path / "ValidationReport.json",
        )
        assert output.exists()

    def test_missing_input_returns_error_code(self):
        from timeline_builder import main

        assert main(["missing.csv"]) == 1

    def test_all_invalid_rows_raises(self, tmp_path):
        bad = tmp_path / "bad.csv"
        bad.write_text(
            "Project,Task Name,Start Date,Finish Date,Critical Date\n"
            ",,,,\n"
            "P,T,,,\n",
            encoding="utf-8",
        )
        with pytest.raises(ValueError, match="No valid tasks"):
            build_timeline(
                bad,
                tmp_path / "out.xlsx",
                CONFIG,
                report_path=tmp_path / "ValidationReport.json",
            )


class TestMonthHeaders:
    def test_month_merges_cover_all_week_columns(self, built_small):
        ws = built_small.active
        first_week_col = len(INFO_COLUMNS) + 1
        merged = [
            (r.min_col, r.max_col)
            for r in ws.merged_cells.ranges
            if r.min_row == 1 and r.min_col >= first_week_col
        ]
        covered = set()
        for start_col, end_col in merged:
            covered.update(range(start_col, end_col + 1))
        expected = set(range(first_week_col, ws.max_column + 1))
        assert covered == expected

    def test_month_labels_are_non_empty(self, built_small):
        ws = built_small.active
        labels = {
            ws.cell(1, r.min_col).value
            for r in ws.merged_cells.ranges
            if r.min_row == 1 and r.min_col > len(INFO_COLUMNS)
        }
        assert labels
        assert all(label for label in labels)

    def test_year_boundary_month_labels(self, tmp_path):
        rows = [
            TaskRow("Y", "Cross year", date(2025, 12, 20), date(2026, 1, 10), None),
        ]
        weeks = build_week_columns(rows)
        categories, critical, default, theme = load_config(CONFIG)
        wb = build_workbook([("Y", rows)], weeks, categories, critical, default, theme)
        ws = wb.active
        labels = set()
        for col in range(len(INFO_COLUMNS) + 1, ws.max_column + 1):
            value = ws.cell(1, col).value
            if value:
                labels.add(value)
        assert "December 2025" in labels
        assert "January 2026" in labels


class TestWeeklyHeaders:
    def test_week_count_matches_span(self):
        tasks = [
            TaskRow("P", "T", date(2026, 1, 6), date(2026, 2, 20), None),
        ]
        weeks = build_week_columns(tasks)
        assert weeks[0].start == week_start(date(2026, 1, 6))
        assert weeks[-1].start == week_start(date(2026, 2, 20))
        assert len(weeks) >= 7

    def test_week_header_row_contains_date_ranges(self, built_small):
        ws = built_small.active
        sample = ws.cell(2, len(INFO_COLUMNS) + 1).value
        assert "\n" in sample
        assert "Jan" in sample or "Feb" in sample

    def test_column_headers_use_w_labels(self, built_small):
        ws = built_small.active
        assert ws.cell(3, len(INFO_COLUMNS) + 1).value == "W1"


class TestProjectGrouping:
    def test_project_count(self, small_input):
        frame = read_schedule(small_input)
        tasks, _ = validate_and_normalize(frame)
        grouped = group_tasks_by_project(tasks)
        assert len(grouped) == 3

    def test_contiguous_blocks_preserve_order(self):
        tasks = [
            TaskRow("B", "t1", date(2026, 1, 1), date(2026, 1, 2), None),
            TaskRow("A", "t2", date(2026, 1, 3), date(2026, 1, 4), None),
            TaskRow("B", "t3", date(2026, 1, 5), date(2026, 1, 6), None),
        ]
        grouped = group_tasks_by_project(tasks)
        assert [name for name, _ in grouped] == ["B", "A", "B"]
        assert [len(block) for _, block in grouped] == [1, 1, 1]

    def test_duplicate_project_names_create_separate_blocks(self, edge_input):
        frame = read_schedule(edge_input)
        tasks, _ = validate_and_normalize(frame)
        grouped = group_tasks_by_project(tasks)
        dup_blocks = [name for name, _ in grouped if name == "Duplicate Project Name"]
        assert len(dup_blocks) == 2

    def test_workbook_has_project_header_rows(self, built_small):
        ws = built_small.active
        headers = []
        for row in range(4, ws.max_row + 1):
            if ws.cell(row, 1).value and not ws.cell(row, 2).value:
                headers.append(ws.cell(row, 1).value)
        assert len(headers) == 3


class TestOutlineCollapse:
    def test_child_rows_have_outline_level_one(self, built_small):
        ws = built_small.active
        child_levels = [
            ws.row_dimensions[row].outlineLevel
            for row in range(4, ws.max_row + 1)
            if ws.cell(row, 2).value
        ]
        assert child_levels
        assert all(level == 1 for level in child_levels)

    def test_summary_below_disabled(self, built_small):
        ws = built_small.active
        assert ws.sheet_properties.outlinePr.summaryBelow is False


class TestColorAssignment:
    def test_known_phase_colors(self, config_bundle):
        categories, _, default, _ = config_bundle
        assert task_style("Concepting", categories, default).fill == "1F3864"
        assert task_style("Shoot day", categories, default).fill == "9DC3E6"
        assert task_style("Random task", categories, default).fill == "D9D9D9"

    def test_review_wins_over_shoot_in_compound_name(self, config_bundle):
        categories, _, default, _ = config_bundle
        style = task_style("Post shoot review", categories, default)
        assert style.fill == "A6A6A6"

    def test_workbook_applies_phase_fill(self, built_small):
        ws = built_small.active
        for row in range(4, ws.max_row + 1):
            name = ws.cell(row, 2).value
            if name and "Concepting" in name:
                col = len(INFO_COLUMNS) + 1
                fill = ws.cell(row, col).fill.fgColor.rgb
                assert fill.endswith("1F3864")
                break
        else:
            pytest.fail("Expected a Concepting row")


class TestCriticalDateOverrides:
    def test_critical_week_uses_red_fill(self, built_small):
        ws = built_small.active
        for row in range(4, ws.max_row + 1):
            critical_value = ws.cell(row, 5).value
            if not critical_value:
                continue
            red_found = False
            for col in range(len(INFO_COLUMNS) + 1, ws.max_column + 1):
                fill = ws.cell(row, col).fill.fgColor.rgb
                if fill and fill.endswith("C00000"):
                    red_found = True
                    break
            assert red_found, f"Row {row} has critical date but no red bar"

    def test_critical_outside_task_range_renders_on_timeline(self):
        task = TaskRow(
            "P",
            "Outside critical",
            date(2026, 5, 1),
            date(2026, 5, 7),
            date(2026, 6, 1),
        )
        weeks = build_week_columns([task])
        assert critical_week_index(task, weeks) is not None
        active = weeks_for_range(task.start_date, task.finish_date, weeks)
        critical = critical_week_index(task, weeks)
        assert critical not in active

    def test_critical_only_milestone_renders_red_bar(self, tmp_path):
        task = TaskRow(
            "P",
            "Outside critical",
            date(2026, 5, 1),
            date(2026, 5, 7),
            date(2026, 6, 1),
        )
        categories, critical, default, theme = load_config(CONFIG)
        weeks = build_week_columns([task])
        wb = build_workbook([("P", [task])], weeks, categories, critical, default, theme)
        ws = wb.active
        task_row = 5  # row 4 is project header
        reds = 0
        for col in range(len(INFO_COLUMNS) + 1, ws.max_column + 1):
            fill = ws.cell(task_row, col).fill.fgColor.rgb
            if fill and fill.endswith("C00000"):
                reds += 1
        assert reds >= 1


class TestTimelineSpan:
    def test_span_includes_critical_dates(self):
        task = TaskRow(
            "P",
            "Outside critical",
            date(2026, 5, 1),
            date(2026, 5, 7),
            date(2026, 6, 1),
        )
        weeks = build_week_columns([task])
        assert weeks[-1].end >= date(2026, 6, 1)

    def test_span_uses_min_start_max_finish(self):
        tasks = [
            TaskRow("P", "A", date(2026, 3, 1), date(2026, 3, 7), None),
            TaskRow("P", "B", date(2026, 1, 6), date(2026, 2, 1), None),
        ]
        weeks = build_week_columns(tasks)
        assert weeks[0].start == week_start(date(2026, 1, 6))
        assert weeks[-1].end >= date(2026, 2, 1)

    def test_one_day_task_covers_single_week(self):
        task = TaskRow("P", "One day", date(2026, 4, 1), date(2026, 4, 1), None)
        weeks = build_week_columns([task])
        active = weeks_for_range(task.start_date, task.finish_date, weeks)
        assert len(active) == 1


class TestFreezePanes:
    def test_freeze_at_first_week_column(self, built_small):
        ws = built_small.active
        assert ws.freeze_panes == "F4"


class TestEdgeCaseNormalization:
    def test_blank_project_is_skipped(self, edge_input):
        frame = read_schedule(edge_input)
        tasks, report = validate_and_normalize(frame)
        assert not any(task.project == "nan" for task in tasks)
        assert not any(task.task_name == "Blank project row" for task in tasks)
        assert any(issue.reason == "missing_project" for issue in report.skipped)

    def test_duplicate_task_names_both_kept(self, edge_input):
        frame = read_schedule(edge_input)
        tasks, _ = validate_and_normalize(frame)
        dupes = [t for t in tasks if t.task_name == "Duplicate task name"]
        assert len(dupes) == 2


class TestTaskLabelPlacement:
    def test_label_only_on_first_active_week(self, built_small):
        ws = built_small.active
        for row in range(4, ws.max_row + 1):
            name = ws.cell(row, 2).value
            if not name:
                continue
            labels = []
            for col in range(len(INFO_COLUMNS) + 1, ws.max_column + 1):
                value = ws.cell(row, col).value
                fill = ws.cell(row, col).fill.fgColor.rgb
                if fill and not fill.endswith(("FFFFFF", "F9F9F9")):
                    labels.append(value)
            if len(labels) > 1:
                assert labels.count(name) <= 1
