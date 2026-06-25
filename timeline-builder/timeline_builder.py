#!/usr/bin/env python3
"""Convert Smartsheet schedule exports into a presentation-quality project roadmap."""

from __future__ import annotations

import argparse
import logging
import sys
from dataclasses import dataclass
from datetime import date, timedelta
from pathlib import Path
from typing import Any

import pandas as pd
import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from color_matcher import task_style
from models import (
    INFO_COLUMNS,
    CategoryRule,
    ColorStyle,
    PresentationTheme,
    TaskRow,
    WeekColumn,
)
from validation import REPORT_FILENAME, ValidationReport, validate_and_normalize

DEFAULT_OUTPUT = "Timeline.xlsx"
DEFAULT_CONFIG = Path(__file__).resolve().parent / "colors.yaml"
DEFAULT_GOOGLE_CONFIG = Path(__file__).resolve().parent / "google_config.yaml"

logger = logging.getLogger(__name__)


@dataclass
class BuildResult:
    output_path: Path | None
    report: ValidationReport
    project_blocks: int = 0
    timeline_weeks: int = 0


def load_config(config_path: Path) -> tuple[
    list[CategoryRule],
    ColorStyle,
    ColorStyle,
    PresentationTheme,
]:
    with config_path.open(encoding="utf-8") as handle:
        raw = yaml.safe_load(handle)

    categories: list[CategoryRule] = []
    for entry in raw.get("categories", []):
        categories.append(
            CategoryRule(
                name=entry["name"],
                patterns=tuple(p.lower() for p in entry.get("patterns", [])),
                style=ColorStyle(
                    fill=_normalize_hex(entry["fill"]),
                    font=_normalize_hex(entry.get("font", "FFFFFF")),
                ),
            )
        )

    critical = raw.get("critical", {})
    default = raw.get("default", {})
    presentation = raw.get("presentation", {})

    return (
        categories,
        ColorStyle(
            fill=_normalize_hex(critical.get("fill", "C00000")),
            font=_normalize_hex(critical.get("font", "FFFFFF")),
        ),
        ColorStyle(
            fill=_normalize_hex(default.get("fill", "D9D9D9")),
            font=_normalize_hex(default.get("font", "404040")),
        ),
        PresentationTheme(
            header_fill=_normalize_hex(presentation.get("header_fill", "203864")),
            header_font=_normalize_hex(presentation.get("header_font", "FFFFFF")),
            project_header_fill=_normalize_hex(
                presentation.get("project_header_fill", "E7E6E6")
            ),
            project_header_font=_normalize_hex(
                presentation.get("project_header_font", "203864")
            ),
            month_header_fill=_normalize_hex(
                presentation.get("month_header_fill", "F2F2F2")
            ),
            month_header_font=_normalize_hex(
                presentation.get("month_header_font", "404040")
            ),
            week_header_fill=_normalize_hex(
                presentation.get("week_header_fill", "FAFAFA")
            ),
            week_header_font=_normalize_hex(
                presentation.get("week_header_font", "595959")
            ),
            grid_color=_normalize_hex(presentation.get("grid_color", "D0D0D0")),
            alt_row_fill=_normalize_hex(presentation.get("alt_row_fill", "F9F9F9")),
        ),
    )


def _normalize_hex(value: str) -> str:
    return value.lstrip("#").upper()


def read_schedule(input_path: Path) -> pd.DataFrame:
    suffix = input_path.suffix.lower()
    if suffix == ".csv":
        frame = pd.read_csv(input_path)
    elif suffix in {".xlsx", ".xlsm", ".xls"}:
        frame = pd.read_excel(input_path)
    else:
        raise ValueError(
            f"Unsupported input format '{suffix}'. Use CSV or XLSX exported from Smartsheet."
        )

    frame.columns = [str(column).strip() for column in frame.columns]
    missing = [column for column in INFO_COLUMNS if column not in frame.columns]
    if missing:
        raise ValueError(
            "Input file is missing required columns: "
            + ", ".join(missing)
            + f". Found: {', '.join(frame.columns)}"
        )

    return frame[INFO_COLUMNS].copy()


def week_start(value: date) -> date:
    return value - timedelta(days=value.weekday())


def timeline_bounds(tasks: list[TaskRow]) -> tuple[date, date]:
    earliest = min(task.start_date for task in tasks)
    latest = max(task.finish_date for task in tasks)
    for task in tasks:
        if task.critical_date is not None:
            earliest = min(earliest, task.critical_date)
            latest = max(latest, task.critical_date)
    return earliest, latest


def build_week_columns(tasks: list[TaskRow]) -> list[WeekColumn]:
    earliest, latest = timeline_bounds(tasks)
    cursor = week_start(earliest)
    last_week = week_start(latest)
    weeks: list[WeekColumn] = []
    while cursor <= last_week:
        end = cursor + timedelta(days=6)
        weeks.append(
            WeekColumn(
                index=len(weeks),
                start=cursor,
                end=end,
                month_key=(cursor.year, cursor.month),
            )
        )
        cursor += timedelta(days=7)
    return weeks


def weeks_for_range(start: date, finish: date, weeks: list[WeekColumn]) -> list[int]:
    active: list[int] = []
    for week in weeks:
        if week.end < start or week.start > finish:
            continue
        active.append(week.index)
    return active


def critical_week_index(task: TaskRow, weeks: list[WeekColumn]) -> int | None:
    if task.critical_date is None:
        return None
    for week in weeks:
        if week.start <= task.critical_date <= week.end:
            return week.index
    return None


def group_tasks_by_project(tasks: list[TaskRow]) -> list[tuple[str, list[TaskRow]]]:
    """Group tasks into contiguous project blocks, preserving source row order."""
    if not tasks:
        return []

    blocks: list[tuple[str, list[TaskRow]]] = []
    current_project = tasks[0].project
    current_tasks = [tasks[0]]

    for task in tasks[1:]:
        if task.project == current_project:
            current_tasks.append(task)
        else:
            blocks.append((current_project, current_tasks))
            current_project = task.project
            current_tasks = [task]

    blocks.append((current_project, current_tasks))
    return blocks


def _fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def _font(hex_color: str, *, bold: bool = False, size: int = 10) -> Font:
    return Font(color=hex_color, bold=bold, size=size, name="Calibri")


def _thin_border(color: str) -> Border:
    side = Side(style="thin", color=color)
    return Border(left=side, right=side, top=side, bottom=side)


def _apply_cell(
    ws: Worksheet,
    row: int,
    column: int,
    value: Any,
    *,
    fill: str | None = None,
    font: Font | None = None,
    alignment: Alignment | None = None,
    border: Border | None = None,
    number_format: str | None = None,
) -> None:
    cell = ws.cell(row=row, column=column, value=value)
    if fill is not None:
        cell.fill = _fill(fill)
    if font is not None:
        cell.font = font
    if alignment is not None:
        cell.alignment = alignment
    if border is not None:
        cell.border = border
    if number_format is not None:
        cell.number_format = number_format


def build_workbook(
    grouped_tasks: list[tuple[str, list[TaskRow]]],
    weeks: list[WeekColumn],
    categories: list[CategoryRule],
    critical_style: ColorStyle,
    default_style: ColorStyle,
    theme: PresentationTheme,
) -> Workbook:
    workbook = Workbook()
    ws = workbook.active
    ws.title = "Timeline"

    ws.sheet_properties.outlinePr.summaryBelow = False
    ws.sheet_properties.outlinePr.applyStyles = True

    first_week_col = len(INFO_COLUMNS) + 1
    last_col = first_week_col + len(weeks) - 1 if weeks else len(INFO_COLUMNS)
    border = _thin_border(theme.grid_color)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    month_row = 1
    week_row = 2
    header_row = 3
    data_start_row = 4

    for column, title in enumerate(INFO_COLUMNS, start=1):
        _apply_cell(
            ws,
            month_row,
            column,
            "",
            fill=theme.month_header_fill,
            font=_font(theme.month_header_font, bold=True, size=11),
            alignment=center,
            border=border,
        )
        _apply_cell(
            ws,
            week_row,
            column,
            "",
            fill=theme.week_header_fill,
            font=_font(theme.week_header_font, bold=True),
            alignment=center,
            border=border,
        )
        _apply_cell(
            ws,
            header_row,
            column,
            title,
            fill=theme.header_fill,
            font=_font(theme.header_font, bold=True, size=11),
            alignment=center,
            border=border,
        )

    if weeks:
        month_spans: list[tuple[int, int, str]] = []
        span_start = first_week_col
        current_key = weeks[0].month_key
        for offset, week in enumerate(weeks):
            column = first_week_col + offset
            if week.month_key != current_key:
                month_spans.append(
                    (
                        span_start,
                        column - 1,
                        _month_label(weeks[span_start - first_week_col].start),
                    )
                )
                span_start = column
                current_key = week.month_key
        month_spans.append(
            (
                span_start,
                last_col,
                _month_label(weeks[span_start - first_week_col].start),
            )
        )

        for start_col, end_col, label in month_spans:
            if start_col == end_col:
                _apply_cell(
                    ws,
                    month_row,
                    start_col,
                    label,
                    fill=theme.month_header_fill,
                    font=_font(theme.month_header_font, bold=True, size=11),
                    alignment=center,
                    border=border,
                )
            else:
                ws.merge_cells(
                    start_row=month_row,
                    start_column=start_col,
                    end_row=month_row,
                    end_column=end_col,
                )
                _apply_cell(
                    ws,
                    month_row,
                    start_col,
                    label,
                    fill=theme.month_header_fill,
                    font=_font(theme.month_header_font, bold=True, size=11),
                    alignment=center,
                    border=border,
                )

        for offset, week in enumerate(weeks):
            column = first_week_col + offset
            week_label = f"{week.start.strftime('%b %d')}\n{week.end.strftime('%b %d')}"
            _apply_cell(
                ws,
                week_row,
                column,
                week_label,
                fill=theme.week_header_fill,
                font=_font(theme.week_header_font, bold=True, size=9),
                alignment=center,
                border=border,
            )
            _apply_cell(
                ws,
                header_row,
                column,
                f"W{offset + 1}",
                fill=theme.header_fill,
                font=_font(theme.header_font, bold=True, size=10),
                alignment=center,
                border=border,
            )

    current_row = data_start_row
    alt_toggle = False

    for project, project_tasks in grouped_tasks:
        _apply_cell(
            ws,
            current_row,
            1,
            project,
            fill=theme.project_header_fill,
            font=_font(theme.project_header_font, bold=True, size=11),
            alignment=left,
            border=border,
        )
        for column in range(2, last_col + 1):
            _apply_cell(
                ws,
                current_row,
                column,
                "",
                fill=theme.project_header_fill,
                font=_font(theme.project_header_font, bold=True),
                alignment=left,
                border=border,
            )
        project_header_row = current_row
        current_row += 1

        for task in project_tasks:
            alt_toggle = not alt_toggle
            row_fill = theme.alt_row_fill if alt_toggle else "FFFFFF"
            style = task_style(task.task_name, categories, default_style)
            active_weeks = weeks_for_range(task.start_date, task.finish_date, weeks)
            critical_index = critical_week_index(task, weeks)
            render_weeks = sorted(set(active_weeks))
            if critical_index is not None and critical_index not in render_weeks:
                render_weeks.append(critical_index)
                render_weeks.sort()
            label_week = active_weeks[0] if active_weeks else critical_index

            _apply_cell(
                ws,
                current_row,
                1,
                "",
                fill=row_fill,
                font=_font("404040"),
                alignment=left,
                border=border,
            )
            _apply_cell(
                ws,
                current_row,
                2,
                task.task_name,
                fill=row_fill,
                font=_font("404040"),
                alignment=left,
                border=border,
            )
            _apply_cell(
                ws,
                current_row,
                3,
                task.start_date,
                fill=row_fill,
                font=_font("404040"),
                alignment=center,
                border=border,
                number_format="mmm d, yyyy",
            )
            _apply_cell(
                ws,
                current_row,
                4,
                task.finish_date,
                fill=row_fill,
                font=_font("404040"),
                alignment=center,
                border=border,
                number_format="mmm d, yyyy",
            )
            _apply_cell(
                ws,
                current_row,
                5,
                task.critical_date if task.critical_date else "",
                fill=row_fill,
                font=_font(
                    "C00000" if task.critical_date else "404040",
                    bold=bool(task.critical_date),
                ),
                alignment=center,
                border=border,
                number_format="mmm d, yyyy",
            )

            for offset, week in enumerate(weeks):
                column = first_week_col + offset
                if offset not in render_weeks:
                    _apply_cell(
                        ws,
                        current_row,
                        column,
                        "",
                        fill=row_fill,
                        alignment=center,
                        border=border,
                    )
                    continue

                is_critical = critical_index is not None and offset == critical_index
                is_duration = offset in active_weeks

                if is_critical:
                    cell_style = critical_style
                elif is_duration:
                    cell_style = style
                else:
                    cell_style = critical_style

                label = ""
                if label_week is not None and offset == label_week:
                    label = task.task_name

                _apply_cell(
                    ws,
                    current_row,
                    column,
                    label,
                    fill=cell_style.fill,
                    font=_font(cell_style.font, bold=bool(label)),
                    alignment=Alignment(
                        horizontal="left" if label else "center",
                        vertical="center",
                        wrap_text=bool(label),
                    ),
                    border=border,
                )

            ws.row_dimensions[current_row].outlineLevel = 1
            current_row += 1

        if current_row > project_header_row + 1:
            ws.row_dimensions.group(
                project_header_row + 1,
                current_row - 1,
                outline_level=1,
                hidden=False,
            )

    ws.freeze_panes = ws.cell(row=data_start_row, column=first_week_col)

    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 14
    ws.column_dimensions["E"].width = 14
    for offset in range(len(weeks)):
        ws.column_dimensions[get_column_letter(first_week_col + offset)].width = 11

    ws.row_dimensions[month_row].height = 22
    ws.row_dimensions[week_row].height = 30
    ws.row_dimensions[header_row].height = 20
    ws.sheet_view.showGridLines = True
    ws.print_title_rows = "1:3"
    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0

    return workbook


def _month_label(value: date) -> str:
    return value.strftime("%B %Y")


def _default_report_path(output_path: Path) -> Path:
    return output_path.parent / REPORT_FILENAME


def print_operator_summary(
    result: BuildResult,
    report_path: Path,
    publish_result: Any | None = None,
) -> None:
    report = result.report
    print("")
    print("Timeline Generation Summary")
    print("---------------------------")
    print(f"Rows Processed:  {report.total_rows}")
    print(f"Rows Included:   {report.valid_rows}")
    print(f"Rows Skipped:    {report.skipped_rows}")
    print(f"Projects:        {result.project_blocks}")
    print(f"Timeline Weeks:  {result.timeline_weeks}")
    print(f"Warnings:        {report.warning_count}")
    print(f"Errors:          {report.error_count}")
    if result.output_path:
        print(f"Timeline:        {result.output_path}")
    else:
        print("Timeline:        not generated")
    print(f"ValidationReport: {report_path}")

    if publish_result is not None:
        print("")
        print(f"Google Publish Status: {publish_result.status}")
        if publish_result.sheet_url:
            print(f"Google Sheet URL:      {publish_result.sheet_url}")
        elif publish_result.error:
            print(f"Google Publish Error:  {publish_result.error}")
        else:
            print("Google Sheet URL:      not available")
        if publish_result.updated_at:
            print(
                "Sheet Updated Time:    "
                f"{publish_result.updated_at.strftime('%Y-%m-%d %H:%M:%S %Z')}"
            )
    print("")


def build_timeline(
    input_path: Path,
    output_path: Path,
    config_path: Path,
    *,
    strict: bool = False,
    report_path: Path | None = None,
) -> BuildResult:
    categories, critical_style, default_style, theme = load_config(config_path)
    frame = read_schedule(input_path)
    tasks, report = validate_and_normalize(frame, source_file=str(input_path))
    report.strict_mode = strict

    if report_path is None:
        report_path = _default_report_path(output_path)

    if not tasks:
        report.save(report_path)
        raise ValueError("No valid tasks found in the input schedule.")

    if strict and report.has_errors():
        report.save(report_path)
        return BuildResult(output_path=None, report=report)

    weeks = build_week_columns(tasks)
    grouped = group_tasks_by_project(tasks)
    workbook = build_workbook(
        grouped,
        weeks,
        categories,
        critical_style,
        default_style,
        theme,
    )
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    report.timeline_generated = True
    report.save(report_path)

    result = BuildResult(
        output_path=output_path,
        report=report,
        project_blocks=len(grouped),
        timeline_weeks=len(weeks),
    )
    logger.info(
        "Wrote %s (%d project blocks, %d tasks, %d weeks).",
        output_path,
        result.project_blocks,
        len(tasks),
        result.timeline_weeks,
    )
    return result


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Convert Smartsheet schedule exports into Timeline.xlsx roadmaps.",
    )
    parser.add_argument(
        "input",
        type=Path,
        help="Smartsheet CSV or XLSX export",
    )
    parser.add_argument(
        "-o",
        "--output",
        type=Path,
        default=Path(DEFAULT_OUTPUT),
        help=f"Output workbook path (default: {DEFAULT_OUTPUT})",
    )
    parser.add_argument(
        "-c",
        "--config",
        type=Path,
        default=DEFAULT_CONFIG,
        help=f"Color mapping YAML config (default: {DEFAULT_CONFIG.name})",
    )
    parser.add_argument(
        "--report",
        type=Path,
        default=None,
        help=f"Validation report path (default: alongside output as {REPORT_FILENAME})",
    )
    parser.add_argument(
        "--strict",
        action="store_true",
        help="Fail without generating timeline when validation errors exist",
    )
    parser.add_argument(
        "--publish-google",
        action="store_true",
        help="Publish Timeline.xlsx to Google Sheets after generation",
    )
    parser.add_argument(
        "--sheet-id",
        type=str,
        default=None,
        help="Existing Google Spreadsheet ID to update (optional)",
    )
    parser.add_argument(
        "--google-config",
        type=Path,
        default=DEFAULT_GOOGLE_CONFIG,
        help=f"Google publishing config (default: {DEFAULT_GOOGLE_CONFIG.name})",
    )
    parser.add_argument(
        "-v",
        "--verbose",
        action="store_true",
        help="Enable verbose logging",
    )
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    logging.basicConfig(
        level=logging.DEBUG if args.verbose else logging.INFO,
        format="%(levelname)s: %(message)s",
    )

    if not args.input.exists():
        logger.error("Input file not found: %s", args.input)
        return 1
    if not args.config.exists():
        logger.error("Config file not found: %s", args.config)
        return 1

    report_path = args.report or _default_report_path(args.output)

    try:
        result = build_timeline(
            args.input,
            args.output,
            args.config,
            strict=args.strict,
            report_path=report_path,
        )
    except ValueError as exc:
        logger.error("%s", exc)
        return 2 if "No valid tasks" in str(exc) else 1
    except Exception as exc:  # noqa: BLE001 - surface user-facing CLI errors
        logger.error("%s", exc)
        return 1

    if result.output_path is None:
        print_operator_summary(result, report_path)
        logger.error(
            "Strict mode: %d validation error(s). Timeline not generated.",
            result.report.error_count,
        )
        return 3

    publish_result = None
    if args.publish_google:
        from google_sheets_publisher import (
            GooglePublishResult,
            load_google_config,
            publish_timeline_xlsx,
        )

        if not args.google_config.exists():
            logger.error("Google config file not found: %s", args.google_config)
            publish_result = GooglePublishResult(
                success=False,
                status="failed",
                error=f"Google config file not found: {args.google_config}",
            )
        else:
            try:
                google_config = load_google_config(args.google_config)
                publish_result = publish_timeline_xlsx(
                    result.output_path,
                    google_config,
                    sheet_id=args.sheet_id,
                    source_name=args.input.stem,
                )
            except Exception as exc:  # noqa: BLE001 - publish must not destroy Excel output
                logger.error("Google Sheets publish failed: %s", exc)
                publish_result = GooglePublishResult(
                    success=False,
                    status="failed",
                    error=str(exc),
                )

    print_operator_summary(result, report_path, publish_result)
    if publish_result and not publish_result.success:
        logger.warning(
            "Timeline.xlsx and ValidationReport.json were written; "
            "Google Sheets publishing failed."
        )
    return 0


if __name__ == "__main__":
    sys.exit(main())
