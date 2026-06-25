"""Publish Timeline.xlsx to Google Sheets as an optional delivery channel."""

from __future__ import annotations

import logging
from dataclasses import dataclass
from datetime import UTC, date, datetime
from pathlib import Path
from typing import Any, Callable

import yaml
from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string
from openpyxl.worksheet.worksheet import Worksheet

logger = logging.getLogger(__name__)

DEFAULT_GOOGLE_CONFIG = Path(__file__).resolve().parent / "google_config.yaml"
SHEETS_SCOPE = "https://www.googleapis.com/auth/spreadsheets"
SHEET_URL_TEMPLATE = "https://docs.google.com/spreadsheets/d/{sheet_id}/edit"


@dataclass(frozen=True)
class GooglePublishOptions:
    tab_name: str = "Timeline"
    create_title_prefix: str = "Timeline Roadmap - "
    preserve_formatting: bool = True
    preserve_merged_cells: bool = True
    preserve_frozen_panes: bool = True
    preserve_row_groups: bool = True


@dataclass(frozen=True)
class GooglePublishConfig:
    credentials_path: Path
    default_sheet_id: str | None
    options: GooglePublishOptions


@dataclass
class GooglePublishResult:
    success: bool
    status: str
    sheet_id: str | None = None
    sheet_url: str | None = None
    updated_at: datetime | None = None
    error: str | None = None
    tab_name: str | None = None


@dataclass
class _SheetPayload:
    values: list[list[Any]]
    merge_ranges: list[tuple[int, int, int, int]]
    repeat_cells: list[dict[str, Any]]
    column_widths: list[tuple[int, int, int]]
    row_groups: list[tuple[int, int]]
    frozen_row_count: int
    frozen_column_count: int


def load_google_config(config_path: Path) -> GooglePublishConfig:
    with config_path.open(encoding="utf-8") as handle:
        raw = yaml.safe_load(handle) or {}

    publish = raw.get("publish", {})
    credentials = raw.get("credentials_path")
    if not credentials:
        raise ValueError("google_config.yaml must set credentials_path.")

    default_sheet_id = raw.get("default_sheet_id")
    if default_sheet_id in (None, "null", ""):
        default_sheet_id = None

    return GooglePublishConfig(
        credentials_path=Path(credentials),
        default_sheet_id=default_sheet_id,
        options=GooglePublishOptions(
            tab_name=str(publish.get("tab_name", "Timeline")),
            create_title_prefix=str(
                publish.get("create_title_prefix", "Timeline Roadmap - ")
            ),
            preserve_formatting=bool(publish.get("preserve_formatting", True)),
            preserve_merged_cells=bool(publish.get("preserve_merged_cells", True)),
            preserve_frozen_panes=bool(publish.get("preserve_frozen_panes", True)),
            preserve_row_groups=bool(publish.get("preserve_row_groups", True)),
        ),
    )


def _hex_to_rgb(hex_color: str) -> dict[str, float]:
    cleaned = hex_color.strip().lstrip("#").upper()
    if len(cleaned) == 8:
        cleaned = cleaned[2:]
    if len(cleaned) != 6:
        raise ValueError(f"Invalid hex color: {hex_color}")
    red = int(cleaned[0:2], 16) / 255
    green = int(cleaned[2:4], 16) / 255
    blue = int(cleaned[4:6], 16) / 255
    return {"red": red, "green": green, "blue": blue}


def _cell_text(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.isoformat()
    return value


def _cell_bg_hex(cell: Cell) -> str | None:
    fill = cell.fill
    if fill is None or fill.fill_type != "solid":
        return None
    color = fill.fgColor
    if color is None or color.type != "rgb" or not color.rgb:
        return None
    rgb = color.rgb.upper()
    if rgb in {"00000000", "00FFFFFF", "FFFFFFFF"}:
        return None
    if len(rgb) == 8:
        return rgb[2:]
    return rgb[-6:]


def _cell_font_style(cell: Cell) -> tuple[str | None, bool, float | None]:
    font = cell.font
    if font is None:
        return None, False, None
    color = None
    if font.color and font.color.type == "rgb" and font.color.rgb:
        rgb = font.color.rgb.upper()
        if len(rgb) == 8:
            color = rgb[2:]
        else:
            color = rgb[-6:]
        if color in {"FFFFFF", "000000"} and not font.bold:
            color = None
    return color, bool(font.bold), font.size


def _style_key(cell: Cell) -> tuple[Any, ...]:
    bg = _cell_bg_hex(cell)
    fg, bold, size = _cell_font_style(cell)
    horizontal = cell.alignment.horizontal if cell.alignment else None
    wrap = cell.alignment.wrap_text if cell.alignment else None
    return (bg, fg, bold, size, horizontal, wrap)


def _repeat_cell_request(
    sheet_id: int,
    start_row: int,
    end_row: int,
    start_col: int,
    end_col: int,
    style_key: tuple[Any, ...],
) -> dict[str, Any] | None:
    bg, fg, bold, size, horizontal, wrap = style_key
    fields: list[str] = []
    user_format: dict[str, Any] = {}

    if bg:
        user_format["backgroundColor"] = _hex_to_rgb(bg)
        fields.append("userEnteredFormat.backgroundColor")

    text_format: dict[str, Any] = {}
    if fg:
        text_format["foregroundColor"] = _hex_to_rgb(fg)
    if bold:
        text_format["bold"] = True
    if size:
        text_format["fontSize"] = size
    if text_format:
        user_format["textFormat"] = text_format
        fields.append("userEnteredFormat.textFormat")

    alignment: dict[str, Any] = {}
    if horizontal:
        google_align = {
            "left": "LEFT",
            "center": "CENTER",
            "right": "RIGHT",
        }.get(str(horizontal).lower(), "LEFT")
        alignment["horizontalAlignment"] = google_align
        fields.append("userEnteredFormat.horizontalAlignment")
    if wrap is not None:
        alignment["wrapStrategy"] = "WRAP" if wrap else "OVERFLOW_CELL"
        fields.append("userEnteredFormat.wrapStrategy")
    if alignment:
        user_format.update(alignment)

    if not fields:
        return None

    return {
        "repeatCell": {
            "range": {
                "sheetId": sheet_id,
                "startRowIndex": start_row,
                "endRowIndex": end_row,
                "startColumnIndex": start_col,
                "endColumnIndex": end_col,
            },
            "cell": {"userEnteredFormat": user_format},
            "fields": ",".join(sorted(set(fields))),
        }
    }


def _parse_freeze_panes(ws: Worksheet) -> tuple[int, int]:
    target = ws.freeze_panes
    if target is None:
        return 0, 0
    if isinstance(target, str):
        column_letters, row_number = coordinate_from_string(target)
        frozen_cols = column_index_from_string(column_letters) - 1
        frozen_rows = int(row_number) - 1
        return frozen_rows, frozen_cols
    frozen_rows = target.row - 1 if target.row else 0
    frozen_cols = target.column - 1 if target.column else 0
    return frozen_rows, frozen_cols


def _collect_row_groups(ws: Worksheet) -> list[tuple[int, int]]:
    groups: list[tuple[int, int]] = []
    row = 1
    max_row = ws.max_row
    while row <= max_row:
        if not _is_project_header_row(ws, row):
            row += 1
            continue
        task_start = row + 1
        row = task_start
        while row <= max_row and ws.row_dimensions[row].outlineLevel == 1:
            row += 1
        task_end = row - 1
        if task_end >= task_start:
            groups.append((task_start - 1, task_end))
    return groups


def _is_project_header_row(ws: Worksheet, row: int) -> bool:
    project = ws.cell(row, 1).value
    task = ws.cell(row, 2).value
    return bool(project) and not task


def _extract_sheet_payload(ws: Worksheet, options: GooglePublishOptions) -> _SheetPayload:
    max_row = ws.max_row
    max_col = ws.max_column
    values: list[list[Any]] = []
    repeat_cells: list[dict[str, Any]] = []

    for row in range(1, max_row + 1):
        row_values: list[Any] = []
        col = 1
        while col <= max_col:
            cell = ws.cell(row, col)
            row_values.append(_cell_text(cell.value))
            col += 1
        values.append(row_values)

    if options.preserve_formatting:
        for row_idx in range(max_row):
            col_idx = 0
            while col_idx < max_col:
                style = _style_key(ws.cell(row_idx + 1, col_idx + 1))
                start_col = col_idx
                while (
                    col_idx + 1 < max_col
                    and _style_key(ws.cell(row_idx + 1, col_idx + 2)) == style
                ):
                    col_idx += 1
                request = _repeat_cell_request(
                    sheet_id=0,
                    start_row=row_idx,
                    end_row=row_idx + 1,
                    start_col=start_col,
                    end_col=col_idx + 1,
                    style_key=style,
                )
                if request is not None:
                    repeat_cells.append(request)
                col_idx += 1

    merge_ranges: list[tuple[int, int, int, int]] = []
    if options.preserve_merged_cells:
        for merged in ws.merged_cells.ranges:
            merge_ranges.append(
                (
                    merged.min_row - 1,
                    merged.max_row,
                    merged.min_col - 1,
                    merged.max_col,
                )
            )

    column_widths: list[tuple[int, int, int]] = []
    for letter, dimension in ws.column_dimensions.items():
        if dimension.width:
            index = column_index_from_string(letter) - 1
            pixels = max(int(dimension.width * 7), 30)
            column_widths.append((index, index + 1, pixels))

    frozen_rows, frozen_cols = _parse_freeze_panes(ws)
    if not options.preserve_frozen_panes:
        frozen_rows, frozen_cols = 0, 0

    row_groups = _collect_row_groups(ws) if options.preserve_row_groups else []

    return _SheetPayload(
        values=values,
        merge_ranges=merge_ranges,
        repeat_cells=repeat_cells,
        column_widths=column_widths,
        row_groups=row_groups,
        frozen_row_count=frozen_rows,
        frozen_column_count=frozen_cols,
    )


def _deep_copy_request(request: dict[str, Any]) -> dict[str, Any]:
    import copy

    return copy.deepcopy(request)


def build_batch_requests(
    payload: _SheetPayload,
    sheet_id: int,
    options: GooglePublishOptions,
) -> list[dict[str, Any]]:
    requests: list[dict[str, Any]] = []

    if options.preserve_merged_cells:
        for start_row, end_row, start_col, end_col in payload.merge_ranges:
            requests.append(
                {
                    "mergeCells": {
                        "range": {
                            "sheetId": sheet_id,
                            "startRowIndex": start_row,
                            "endRowIndex": end_row,
                            "startColumnIndex": start_col,
                            "endColumnIndex": end_col,
                        },
                        "mergeType": "MERGE_ALL",
                    }
                }
            )

    if options.preserve_formatting:
        formatting_requests = []
        for request in payload.repeat_cells:
            cloned = _deep_copy_request(request)
            cloned["repeatCell"]["range"]["sheetId"] = sheet_id
            formatting_requests.append(cloned)
        requests.extend(formatting_requests)

    grid_properties: dict[str, Any] = {}
    if payload.frozen_row_count:
        grid_properties["frozenRowCount"] = payload.frozen_row_count
    if payload.frozen_column_count:
        grid_properties["frozenColumnCount"] = payload.frozen_column_count

    if grid_properties:
        requests.append(
            {
                "updateSheetProperties": {
                    "properties": {
                        "sheetId": sheet_id,
                        "gridProperties": grid_properties,
                    },
                    "fields": ",".join(
                        f"gridProperties.{key}" for key in grid_properties
                    ),
                }
            }
        )

    for start_col, end_col, pixels in payload.column_widths:
        requests.append(
            {
                "updateDimensionProperties": {
                    "range": {
                        "sheetId": sheet_id,
                        "dimension": "COLUMNS",
                        "startIndex": start_col,
                        "endIndex": end_col,
                    },
                    "properties": {"pixelSize": pixels},
                    "fields": "pixelSize",
                }
            }
        )

    if options.preserve_row_groups:
        for start_row, end_row in payload.row_groups:
            requests.append(
                {
                    "addDimensionGroup": {
                        "range": {
                            "sheetId": sheet_id,
                            "dimension": "ROWS",
                            "startIndex": start_row,
                            "endIndex": end_row,
                        }
                    }
                }
            )

    return requests


def _build_sheets_service(credentials_path: Path) -> Any:
    from google.oauth2.service_account import Credentials
    from googleapiclient.discovery import build

    credentials = Credentials.from_service_account_file(
        str(credentials_path),
        scopes=[SHEETS_SCOPE],
    )
    return build("sheets", "v4", credentials=credentials, cache_discovery=False)


def _find_sheet_id(spreadsheet: dict[str, Any], tab_name: str) -> int | None:
    for sheet in spreadsheet.get("sheets", []):
        properties = sheet.get("properties", {})
        if properties.get("title") == tab_name:
            return properties.get("sheetId")
    return None


def _replace_timeline_tab(
    service: Any,
    spreadsheet_id: str,
    tab_name: str,
) -> int:
    spreadsheet = (
        service.spreadsheets()
        .get(spreadsheetId=spreadsheet_id, fields="sheets.properties")
        .execute()
    )
    existing_id = _find_sheet_id(spreadsheet, tab_name)
    requests: list[dict[str, Any]] = []
    if existing_id is not None:
        requests.append({"deleteSheet": {"sheetId": existing_id}})
    requests.append({"addSheet": {"properties": {"title": tab_name}}})
    response = (
        service.spreadsheets()
        .batchUpdate(spreadsheetId=spreadsheet_id, body={"requests": requests})
        .execute()
    )
    for reply in response.get("replies", []):
        if "addSheet" in reply:
            return reply["addSheet"]["properties"]["sheetId"]
    refreshed = (
        service.spreadsheets()
        .get(spreadsheetId=spreadsheet_id, fields="sheets.properties")
        .execute()
    )
    sheet_id = _find_sheet_id(refreshed, tab_name)
    if sheet_id is None:
        raise RuntimeError(f"Failed to create Google Sheets tab '{tab_name}'.")
    return sheet_id


def _create_spreadsheet(
    service: Any,
    title: str,
    tab_name: str,
) -> tuple[str, int]:
    body = {
        "properties": {"title": title},
        "sheets": [{"properties": {"title": tab_name}}],
    }
    response = service.spreadsheets().create(body=body).execute()
    spreadsheet_id = response["spreadsheetId"]
    sheet_id = response["sheets"][0]["properties"]["sheetId"]
    return spreadsheet_id, sheet_id


def publish_timeline_xlsx(
    xlsx_path: Path,
    config: GooglePublishConfig,
    *,
    sheet_id: str | None = None,
    source_name: str | None = None,
    service_factory: Callable[[Path], Any] | None = None,
) -> GooglePublishResult:
    tab_name = config.options.tab_name
    updated_at = datetime.now(UTC)

    try:
        if not xlsx_path.exists():
            raise FileNotFoundError(f"Timeline workbook not found: {xlsx_path}")
        if not config.credentials_path.exists():
            raise FileNotFoundError(
                f"Google credentials not found: {config.credentials_path}"
            )

        workbook = load_workbook(xlsx_path, data_only=True)
        if "Timeline" not in workbook.sheetnames:
            raise ValueError("Timeline.xlsx does not contain a 'Timeline' worksheet.")
        worksheet = workbook["Timeline"]
        payload = _extract_sheet_payload(worksheet, config.options)

        factory = service_factory or _build_sheets_service
        service = factory(config.credentials_path)

        target_sheet_id = sheet_id or config.default_sheet_id
        if target_sheet_id:
            tab_sheet_id = _replace_timeline_tab(service, target_sheet_id, tab_name)
            spreadsheet_id = target_sheet_id
        else:
            title = f"{config.options.create_title_prefix}{source_name or xlsx_path.stem}"
            spreadsheet_id, tab_sheet_id = _create_spreadsheet(service, title, tab_name)

        service.spreadsheets().values().update(
            spreadsheetId=spreadsheet_id,
            range=f"{tab_name}!A1",
            valueInputOption="USER_ENTERED",
            body={"values": payload.values},
        ).execute()

        batch_requests = build_batch_requests(payload, tab_sheet_id, config.options)
        if batch_requests:
            _chunk_batch_update(service, spreadsheet_id, batch_requests)

        url = SHEET_URL_TEMPLATE.format(sheet_id=spreadsheet_id)
        logger.info("Published timeline to Google Sheet %s", url)
        return GooglePublishResult(
            success=True,
            status="published",
            sheet_id=spreadsheet_id,
            sheet_url=url,
            updated_at=updated_at,
            tab_name=tab_name,
        )
    except Exception as exc:  # noqa: BLE001 - publish failures must not crash pipeline
        logger.error("Google Sheets publish failed: %s", exc)
        return GooglePublishResult(
            success=False,
            status="failed",
            updated_at=updated_at,
            error=str(exc),
            tab_name=tab_name,
        )


def _chunk_batch_update(
    service: Any,
    spreadsheet_id: str,
    requests: list[dict[str, Any]],
    chunk_size: int = 500,
) -> None:
    for start in range(0, len(requests), chunk_size):
        chunk = requests[start : start + chunk_size]
        service.spreadsheets().batchUpdate(
            spreadsheetId=spreadsheet_id,
            body={"requests": chunk},
        ).execute()
